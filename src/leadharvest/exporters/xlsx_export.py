"""XLSX export: text cells (phones stay +63…), bold frozen header, About sheet with ODbL note."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, TYPE_STRING
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from leadharvest.exporters.base import (
    MANAGED_COLUMNS,
    NUMERIC_COLUMNS,
    OSM_ATTRIBUTION,
    export_basename,
    lead_to_row,
)
from leadharvest.models import ExportResult, Lead, Run, utcnow_iso

_MAX_WIDTH = 50


class XlsxExporter:
    name = "xlsx"

    def __init__(self, export_dir: Path) -> None:
        self.export_dir = export_dir

    def export(self, leads: list[Lead], run: Run) -> ExportResult:
        return self._write(leads, run, f"{export_basename(run)}.xlsx", "Leads")

    def export_new(self, leads: list[Lead], run: Run) -> ExportResult:
        """Monitoring: only the leads that are new since the previous run."""
        return self._write(leads, run, f"{export_basename(run)}-new.xlsx", "New since last run")

    def _write(self, leads: list[Lead], run: Run, filename: str, sheet: str) -> ExportResult:
        self.export_dir.mkdir(parents=True, exist_ok=True)
        path = self.export_dir / filename
        wb = Workbook()
        ws = wb.active
        ws.title = sheet
        ws.append(list(MANAGED_COLUMNS))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        widths = [len(c) for c in MANAGED_COLUMNS]
        for r, lead in enumerate(leads, start=2):
            row = lead_to_row(lead, run.category)
            for c, column in enumerate(MANAGED_COLUMNS, start=1):
                value = row[column]
                if value is None or value == "":
                    continue
                cell = ws.cell(row=r, column=c)
                if column in NUMERIC_COLUMNS:
                    cell.value = value
                else:
                    # Explicit string type: Excel never turns '+639…' or '=…' into a number/formula.
                    cell.value = ILLEGAL_CHARACTERS_RE.sub("", str(value))
                    cell.data_type = TYPE_STRING
                widths[c - 1] = max(widths[c - 1], len(str(value)))
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = min(width + 2, _MAX_WIDTH)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        about = wb.create_sheet("About")
        for row in (
            ("Attribution", OSM_ATTRIBUTION),
            ("Category", run.category),
            ("Location", run.location),
            ("Run ID", run.id),
            ("Exported at (UTC)", utcnow_iso()),
            ("Rows", str(len(leads))),
        ):
            about.append(row)
        about.column_dimensions["A"].width = 20
        about.column_dimensions["B"].width = 90
        wb.save(path)
        return ExportResult(exporter=self.name, target=str(path), rows_appended=len(leads))
