"""CSV/XLSX output and the Google Sheets upsert against a fake worksheet (Phase 5 gate)."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

import pytest
from gspread.utils import a1_range_to_grid_range
from openpyxl import load_workbook

from leadharvest.exporters.base import MANAGED_COLUMNS, ExportError
from leadharvest.exporters.csv_export import CsvExporter, escape_cell
from leadharvest.exporters.sheets_export import SheetsExporter, plan_upsert, tab_title
from leadharvest.exporters.xlsx_export import XlsxExporter
from leadharvest.models import Lead, Run, utcnow_iso


def run() -> Run:
    return Run(
        id="abcdef12-0000",
        category="dentist",
        location="Makati, Philippines",
        area_name="Makati",
        created_at=utcnow_iso(),
    )


def lead(lead_id: str, name: str = "Smile Dental", **kw: Any) -> Lead:
    now = utcnow_iso()
    return Lead(
        lead_id=lead_id,
        business_name=name,
        name_key=name.lower(),
        phone="+639171234567",
        first_seen_run_id="r",
        first_seen_at=now,
        last_seen_at=now,
        updated_at=now,
        **kw,
    )


# ---- CSV / XLSX -------------------------------------------------------------------------------


def test_escape_cell() -> None:
    assert escape_cell("business_name", '=HYPERLINK("x")') == '\'=HYPERLINK("x")'
    assert escape_cell("business_name", "@SUM(1)") == "'@SUM(1)"
    assert escape_cell("phone", "+639171234567") == "+639171234567"
    assert escape_cell("lon", -121.0) == "-121.0"
    assert escape_cell("email", None) == ""


def test_csv_export_bom_and_injection(tmp_path) -> None:
    leads = [
        lead("L-1", "=cmd|' /C calc'!A0", phones_extra=["+63281234567"]),
        lead("L-2", "Niño Dental"),
    ]
    result = CsvExporter(tmp_path).export(leads, run())
    path = Path(result.target)
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(path.read_text(encoding="utf-8-sig").splitlines()))
    assert rows[0] == list(MANAGED_COLUMNS)
    assert rows[1][1].startswith("'=")
    assert rows[1][3] == "+639171234567"
    assert rows[2][1] == "Niño Dental"
    about = (tmp_path / "dentist-makati-abcdef12.about.txt").read_text(encoding="utf-8")
    assert "OpenStreetMap contributors" in about


def test_xlsx_export_text_phones_and_about_sheet(tmp_path) -> None:
    leads = [lead("L-1", "=1+1", lat=14.5, lon=121.0, flags=["no_https"])]
    result = XlsxExporter(tmp_path).export(leads, run())
    wb = load_workbook(result.target)
    ws = wb["Leads"]
    header = [c.value for c in ws[1]]
    assert header == list(MANAGED_COLUMNS)
    row = {h: c for h, c in zip(header, ws[2], strict=True)}
    assert row["phone"].value == "+639171234567" and row["phone"].data_type == "s"
    assert row["business_name"].value == "=1+1" and row["business_name"].data_type == "s"
    assert row["lat"].value == 14.5
    assert row["flags"].value == "no_https"
    assert ws.freeze_panes == "A2"
    assert "OpenStreetMap" in wb["About"]["B1"].value


# ---- Sheets -----------------------------------------------------------------------------------


class FakeWorksheet:
    """Mimics the gspread Worksheet calls the exporter uses, applying RAW writes to a grid."""

    def __init__(self, grid: list[list[Any]] | None = None, rows: int = 100, cols: int = 26):
        self.title = "dentist - Makati"
        self.grid = [list(r) for r in (grid or [])]
        self.row_count = rows
        self.col_count = cols
        self.calls: list[str] = []
        self.input_options: list[str] = []

    def get_all_values(self) -> list[list[str]]:
        self.calls.append("get_all_values")
        width = max((len(r) for r in self.grid), default=0)
        return [[str(v) for v in r] + [""] * (width - len(r)) for r in self.grid]

    def batch_update(self, data: list[dict[str, Any]], **kwargs: Any) -> None:
        self.calls.append("batch_update")
        self.input_options.append(kwargs.get("value_input_option"))
        for item in data:
            rng = a1_range_to_grid_range(item["range"])
            r0, c0 = rng["startRowIndex"], rng["startColumnIndex"]
            assert r0 < self.row_count and c0 < self.col_count, "write outside the grid"
            for dr, values in enumerate(item["values"]):
                for dc, value in enumerate(values):
                    self._set(r0 + dr, c0 + dc, value)

    def _set(self, r: int, c: int, value: Any) -> None:
        while len(self.grid) <= r:
            self.grid.append([])
        row = self.grid[r]
        while len(row) <= c:
            row.append("")
        row[c] = value

    def resize(self, rows: int | None = None, cols: int | None = None) -> None:
        self.row_count = rows or self.row_count
        self.col_count = cols or self.col_count

    def freeze(self, rows: int | None = None, cols: int | None = None) -> None:
        self.calls.append("freeze")

    def format(self, ranges: str, cell_format: dict[str, Any]) -> None:
        self.calls.append("format")

    def set_basic_filter(self, name: str | None = None) -> None:
        self.calls.append("filter")

    def column(self, header: str) -> list[Any]:
        idx = [str(h) for h in self.grid[0]].index(header)
        return [r[idx] if idx < len(r) else "" for r in self.grid[1:]]


def export_to(ws: FakeWorksheet, leads: list[Lead]):
    exporter = SheetsExporter(lambda title, n: ws, sleep=lambda s: None)
    return exporter.export(leads, run())


def test_new_sheet_gets_header_and_rows() -> None:
    ws = FakeWorksheet()
    result = export_to(ws, [lead("L-1"), lead("L-2", "Bright")])
    assert [str(h) for h in ws.grid[0]] == list(MANAGED_COLUMNS)
    assert ws.column("lead_id") == ["L-1", "L-2"]
    assert ws.column("phone") == ["+639171234567", "+639171234567"]
    assert (result.rows_appended, result.rows_updated) == (2, 0)
    assert set(ws.input_options) == {"RAW"}
    assert ws.calls.count("get_all_values") == 1 and ws.calls.count("batch_update") == 1


def test_reexport_updates_appends_and_keeps_client_notes() -> None:
    ws = FakeWorksheet()
    export_to(ws, [lead("L-1"), lead("L-2", "Bright")])
    ws._set(0, len(MANAGED_COLUMNS), "Notes")
    ws._set(1, len(MANAGED_COLUMNS), "called, wants quote")
    result = export_to(ws, [lead("L-2", "Bright Smile Dental"), lead("L-3", "New Clinic")])
    assert (result.rows_updated, result.rows_appended) == (1, 1)
    assert ws.column("Notes") == ["called, wants quote", "", ""]
    assert ws.column("business_name") == ["Smile Dental", "Bright Smile Dental", "New Clinic"]


def test_client_reorders_columns_and_inserts_own_column() -> None:
    ws = FakeWorksheet()
    export_to(ws, [lead("L-1"), lead("L-2", "Bright")])
    # Client moves 'email' to the front and inserts a 'Status' column in the middle.
    header = [str(h) for h in ws.grid[0]]
    reordered = ["email", "Status", *[h for h in header if h != "email"]]
    new_grid = []
    for r, row in enumerate(ws.grid):
        values = dict(zip(header, row, strict=False))
        status = "Status" if r == 0 else f"st-{r}"
        new_grid.append(
            [values.get("email", ""), status, *[values.get(h, "") for h in reordered[2:]]]
        )
    new_grid[0][0] = "email"
    ws.grid = new_grid
    export_to(ws, [lead("L-1", email="info@smile.ph"), lead("L-2", "Bright")])
    assert ws.column("Status") == ["st-1", "st-2"]
    assert ws.column("email") == ["info@smile.ph", ""]
    assert ws.column("lead_id") == ["L-1", "L-2"]


def test_missing_managed_column_is_appended_after_client_columns() -> None:
    old_managed = [c for c in MANAGED_COLUMNS if c != "tiktok"]
    grid = [[*old_managed, "Owner"], ["L-1", "Old Name", *[""] * (len(old_managed) - 2), "Ana"]]
    ws = FakeWorksheet(grid, cols=len(old_managed) + 1)
    export_to(ws, [lead("L-1", tiktok="https://www.tiktok.com/@smile")])
    header = [str(h) for h in ws.grid[0]]
    assert header.index("tiktok") == len(old_managed) + 1  # after 'Owner'
    assert ws.column("Owner") == ["Ana"]
    assert ws.column("tiktok") == ["https://www.tiktok.com/@smile"]
    assert ws.col_count >= len(header)


def test_formula_text_is_written_raw() -> None:
    ws = FakeWorksheet()
    export_to(ws, [lead("L-1", '=HYPERLINK("http://evil","x")')])
    assert ws.column("business_name") == ['=HYPERLINK("http://evil","x")']
    assert ws.input_options == ["RAW"]


def test_sheet_grows_when_needed() -> None:
    ws = FakeWorksheet(rows=2, cols=5)
    export_to(ws, [lead(f"L-{i}") for i in range(5)])
    assert ws.row_count >= 6
    assert len(ws.column("lead_id")) == 5


def test_non_lead_rows_are_not_rewritten() -> None:
    grid = [list(MANAGED_COLUMNS), ["L-1", "A"], ["", "client summary row"], ["L-2", "B"]]
    ws = FakeWorksheet(grid)
    plan = plan_upsert(ws.get_all_values(), [{"lead_id": "L-1"}, {"lead_id": "L-2"}])
    ranges = [u["range"] for u in plan.updates]
    assert all(not r.startswith("A3") for r in ranges)
    assert plan.updated == 2


def test_retry_on_429_then_error() -> None:
    class Resp:
        status_code = 429

    class ApiError(Exception):
        response = Resp()

    attempts = {"n": 0}

    def flaky_open(title: str, n: int):
        attempts["n"] += 1
        raise ApiError("quota")

    with pytest.raises(ExportError):
        SheetsExporter(flaky_open, sleep=lambda s: None, attempts=3).export([lead("L-1")], run())
    assert attempts["n"] == 3


def test_tab_title() -> None:
    assert tab_title("dentist", "Makati") == "dentist - Makati"
    assert tab_title("x", "a/b:c[d]") == "x - a b c d"
    assert len(tab_title("x", "y" * 300)) == 100
